from datetime import date
from decimal import Decimal
from typing import Optional, List

from fastapi import APIRouter, HTTPException, Query, UploadFile, File
from fastapi.responses import StreamingResponse, JSONResponse

from app.database import get_raw_connection

import csv
import io
import json
import xml.etree.ElementTree as ET

from docx import Document
from openpyxl import Workbook
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


router = APIRouter(prefix="/admin/reports", tags=["admin-reports"])


def _serialize_value(value):
    if isinstance(value, Decimal):
        return float(value)
    if hasattr(value, "isoformat"):
        # datetime / date / time
        return value.isoformat()
    return value


def _serialize_row(row: dict) -> dict:
    return {k: _serialize_value(v) for k, v in row.items()}


def fetch_orders(
    date_from: Optional[date] = None,
    date_to: Optional[date] = None,
    status: Optional[str] = None,
) -> List[dict]:
    """
    Загружает историю заказов из таблицы appointments с привязкой к клиенту и услуге.
    ЧИСТЫЙ SQL (без ORM) – для демонстрации работы с БД в курсовом проекте.
    """
    sql = """
        SELECT
            a.id,
            a.id as order_number,
            a.date as order_date,
            CASE WHEN a.status = 'completed' THEN a.date ELSE NULL END as completion_date,
            CASE
                WHEN a.status = 'pending' THEN 'created'
                WHEN a.status = 'confirmed' THEN 'in_progress'
                WHEN a.status = 'completed' THEN 'completed'
                WHEN a.status = 'cancelled' THEN 'cancelled'
                ELSE 'created'
            END as order_status,
            CASE
                WHEN a.status = 'completed' THEN 'paid'
                ELSE 'pending'
            END as payment_status,
            a.price as total_amount,
            a.client_name,
            a.client_email,
            a.service_name
        FROM appointments a
        WHERE 1=1
    """
    params: list = []

    if date_from:
        sql += " AND a.date >= %s"
        params.append(date_from)
    if date_to:
        sql += " AND a.date <= %s"
        params.append(date_to)
    if status and status != "all":
        sql += " AND CASE"
        sql += " WHEN a.status = 'pending' THEN 'created'"
        sql += " WHEN a.status = 'confirmed' THEN 'in_progress'"
        sql += " WHEN a.status = 'completed' THEN 'completed'"
        sql += " WHEN a.status = 'cancelled' THEN 'cancelled'"
        sql += " ELSE 'created' END = %s"
        params.append(status)

    sql += " ORDER BY a.date DESC, a.id DESC LIMIT 500"

    with get_raw_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(sql, params)
            rows = cur.fetchall()
    return [_serialize_row(dict(r)) for r in rows]


@router.get("/orders/history")
def get_orders_history(
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    status: Optional[str] = Query("all"),
):
    """
    История заказов для админ‑панели (таблица на сайте).
    """
    data = fetch_orders(date_from=date_from, date_to=date_to, status=status)
    return JSONResponse(content=data)


@router.get("/orders/export")
def export_orders(
    format: str = Query("json", pattern="^(json|csv|xml)$"),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    status: Optional[str] = Query("all"),
):
    """
    Экспорт истории заказов в JSON / CSV / XML.
    """
    orders = fetch_orders(date_from=date_from, date_to=date_to, status=status)

    if format == "json":
        return JSONResponse(
            content=orders,
            media_type="application/json",
        )

    if format == "csv":
        if not orders:
            fieldnames = [
                "id",
                "order_number",
                "order_date",
                "completion_date",
                "order_status",
                "payment_status",
                "total_amount",
                "client_name",
                "client_email",
                "service_name",
            ]
        else:
            fieldnames = list(orders[0].keys())

        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
        writer.writeheader()
        for row in orders:
            writer.writerow(row)

        data = output.getvalue().encode("utf-8-sig")
        return StreamingResponse(
            io.BytesIO(data),
            media_type="text/csv",
            headers={"Content-Disposition": 'attachment; filename="orders_export.csv"'},
        )

    if format == "xml":
        root = ET.Element("orders")
        for row in orders:
            order_el = ET.SubElement(root, "order")
            for key, value in row.items():
                child = ET.SubElement(order_el, key)
                child.text = "" if value is None else str(value)

        xml_bytes = ET.tostring(root, encoding="utf-8", xml_declaration=True)
        return StreamingResponse(
            io.BytesIO(xml_bytes),
            media_type="application/xml",
            headers={"Content-Disposition": 'attachment; filename="orders_export.xml"'},
        )

    raise HTTPException(status_code=400, detail="Неподдерживаемый формат")


@router.post("/orders/import")
async def import_orders(
    format: str = Query("json", pattern="^(json|csv|xml)$"),
    file: UploadFile = File(...),
):
    """
    Пример импорта заказов из JSON/CSV/XML в таблицу orders.
    Для учебных целей: простая вставка без сложной валидации.
    """
    content = await file.read()

    try:
        with get_raw_connection() as conn:
            with conn.cursor() as cur:
                if format == "json":
                    items = json.loads(content.decode("utf-8"))
                elif format == "csv":
                    text = content.decode("utf-8-sig")
                    reader = csv.DictReader(io.StringIO(text), delimiter=";")
                    items = list(reader)
                else:  # xml
                    root = ET.fromstring(content)
                    items = []
                    for order_el in root.findall("order"):
                        row = {}
                        for child in list(order_el):
                            row[child.tag] = child.text
                        items.append(row)

                inserted = 0
                for row in items:
                    # В учебных целях заполняем только часть полей.
                    # Номер заказа можно оставить NULL — сгенерируется триггером.
                    sql = """
                        INSERT INTO orders (
                            order_date,
                            completion_date,
                            order_status,
                            payment_status,
                            total_amount
                        )
                        VALUES (%s, %s, %s, %s, %s)
                    """
                    cur.execute(
                        sql,
                        (
                            row.get("order_date"),
                            row.get("completion_date"),
                            row.get("order_status", "created"),
                            row.get("payment_status", "pending"),
                            row.get("total_amount") or 0,
                        ),
                    )
                    inserted += 1

                conn.commit()
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Ошибка при импорте данных: {exc}")

    return {"status": "ok", "inserted": inserted}


def _build_orders_doc(orders: List[dict]) -> bytes:
    doc = Document()
    doc.add_heading("История заказов детейлинг‑центра", level=1)

    if not orders:
        doc.add_paragraph("Данных по заказам нет.")
    else:
        table = doc.add_table(rows=1, cols=6)
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = "Номер"
        hdr_cells[1].text = "Дата"
        hdr_cells[2].text = "Клиент"
        hdr_cells[3].text = "Услуга"
        hdr_cells[4].text = "Статус"
        hdr_cells[5].text = "Сумма"

        for row in orders:
            cells = table.add_row().cells
            cells[0].text = str(row.get("order_number") or "")
            cells[1].text = str(row.get("order_date") or "")
            cells[2].text = str(row.get("client_name") or "")
            cells[3].text = str(row.get("service_name") or "")
            cells[4].text = str(row.get("order_status") or "")
            cells[5].text = str(row.get("total_amount") or "")

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _build_orders_excel(orders: List[dict]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "История заказов"

    headers = [
        "Номер",
        "Дата",
        "Клиент",
        "Услуга",
        "Статус",
        "Оплата",
        "Сумма",
    ]
    ws.append(headers)

    for row in orders:
        ws.append(
            [
                row.get("order_number"),
                row.get("order_date"),
                row.get("client_name"),
                row.get("service_name"),
                row.get("order_status"),
                row.get("payment_status"),
                row.get("total_amount"),
            ]
        )

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def _build_orders_pdf(orders: List[dict]) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    width, height = A4

    c.setFont("Helvetica-Bold", 16)
    c.drawString(40, height - 50, "История заказов детейлинг‑центра")

    y = height - 90
    c.setFont("Helvetica", 10)
    headers = ["Номер", "Дата", "Клиент", "Услуга", "Статус", "Сумма"]
    col_widths = [70, 70, 120, 120, 70, 60]

    # Заголовок таблицы
    x = 40
    for i, h in enumerate(headers):
        c.drawString(x, y, h)
        x += col_widths[i]

    y -= 15

    for row in orders:
        if y < 60:
            c.showPage()
            y = height - 60
            c.setFont("Helvetica", 10)

        x = 40
        values = [
            str(row.get("order_number") or ""),
            str(row.get("order_date") or ""),
            str(row.get("client_name") or ""),
            str(row.get("service_name") or ""),
            str(row.get("order_status") or ""),
            str(row.get("total_amount") or ""),
        ]
        for i, val in enumerate(values):
            c.drawString(x, y, val[:20])
            x += col_widths[i]
        y -= 14

    c.showPage()
    c.save()
    buf.seek(0)
    return buf.getvalue()


@router.get("/orders/report")
def download_orders_report(
    format: str = Query("word", pattern="^(word|excel|pdf)$"),
    date_from: Optional[date] = Query(None),
    date_to: Optional[date] = Query(None),
    status: Optional[str] = Query("all"),
):
    """
    Отчет «История заказов» в форматах Word / Excel / PDF.
    """
    orders = fetch_orders(date_from=date_from, date_to=date_to, status=status)

    if format == "word":
        data = _build_orders_doc(orders)
        filename = "orders_report.docx"
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif format == "excel":
        data = _build_orders_excel(orders)
        filename = "orders_report.xlsx"
        media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    else:  # pdf
        data = _build_orders_pdf(orders)
        filename = "orders_report.pdf"
        media_type = "application/pdf"

    return StreamingResponse(
        io.BytesIO(data),
        media_type=media_type,
        headers={"Content-Disposition": f'attachment; filename="{filename}"'},
    )


@router.get("/clients/export-csv")
def export_clients_csv():
    """
    Экспорт клиентов (users с ролью 'client') в CSV.
    Формат: full_name;email;phone;created_at
    """
    sql = """
        SELECT
            full_name,
            email,
            phone,
            created_at
        FROM users
        WHERE role = 'client'
        ORDER BY created_at
    """
    with get_raw_connection() as conn:
        with conn.cursor() as cur:
            cur.execute(sql)
            rows = [dict(r) for r in cur.fetchall()]
    rows = [_serialize_row(r) for r in rows]

    fieldnames = ["full_name", "email", "phone", "created_at"]
    output = io.StringIO()
    writer = csv.DictWriter(output, fieldnames=fieldnames, delimiter=";")
    writer.writeheader()
    for row in rows:
        writer.writerow(
            {
                "full_name": row.get("full_name") or "",
                "email": row.get("email") or "",
                "phone": row.get("phone") or "",
                "created_at": row.get("created_at") or "",
            }
        )

    data = output.getvalue().encode("utf-8-sig")
    return StreamingResponse(
        io.BytesIO(data),
        media_type="text/csv",
        headers={"Content-Disposition": 'attachment; filename="clients_export.csv"'},
    )


@router.post("/clients/import-csv")
async def import_clients_csv(file: UploadFile = File(...)):
    """
    Импорт клиентов из CSV в таблицу users (role='client').
    Ожидаемые колонки: full_name;email;phone
    Если пользователь с таким email уже существует, новая запись не создаётся.
    """
    content = await file.read()
    text = content.decode("utf-8-sig")
    reader = csv.DictReader(io.StringIO(text), delimiter=";")
    items = list(reader)

    try:
        with get_raw_connection() as conn:
            with conn.cursor() as cur:
                inserted = 0
                for row in items:
                    email = (row.get("email") or "").strip()
                    full_name = (row.get("full_name") or "").strip()
                    phone = (row.get("phone") or "").strip()

                    if not email:
                        # Пропускаем строки без email
                        continue

                    # Проверяем, существует ли уже такой пользователь
                    cur.execute(
                        "SELECT id FROM users WHERE email = %s",
                        (email,),
                    )
                    existing = cur.fetchone()
                    if existing:
                        continue

                    # Создаем нового клиента
                    cur.execute(
                        """
                        INSERT INTO users (email, password_hash, full_name, phone, role, is_active)
                        VALUES (%s, %s, %s, %s, 'client', TRUE)
                        """,
                        (
                            email,
                            "imported",  # тех. значение для курсового проекта
                            full_name or email,
                            phone,
                        ),
                    )
                    inserted += 1

                conn.commit()
    except Exception as exc:  # noqa: BLE001
        raise HTTPException(status_code=400, detail=f"Ошибка при импорте клиентов: {exc}")

    return {"status": "ok", "inserted": inserted}


